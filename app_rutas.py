import streamlit as st
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

# Configuración visual
st.set_page_config(page_title="Gestión de Rutas y Calidad Res. 3100", page_icon="♻️", layout="wide")

st.title("♻️ Plataforma de Gestión de Rutas y Calidad Res. 3100")
st.markdown("Automatización de cruces, novedades RF, trazabilidad de crónicos e indicadores de la **Resolución 3100 de 2019**.")

# ==========================================
# 1. CUADRO OFICIAL DE NOVEDADES RF (1 A 15)
# ==========================================
CODIGOS_NOVEDADES_RF = {
    1: "HOSPITALIZADO", 2: "NO SE EN CUENTRA EN EL DOMICILIO", 3: "CITA MEDICA",
    4: "SIN SERVICIO", 5: "NO GENERO RESIDUOS", 6: "BOTO LOS RESIDUOS A LA BASURA",
    7: "FALLECIDO", 8: "DAÑOS DEL VEHICULO", 9: "FACTORES AMBIENTALES",
    10: "LABORES ADMINISTRATIVAS", 11: "ALTERACION ORDEN PUBLICO",
    12: "NO REQUIERE RECOLECCION", 13: "SERVICIOS SUSPENDIDOS TEMPORALMENTE",
    14: "RECOLECCION CADA 15 DIAS", 15: "SE CAMBIO DOMICILIO"
}

MAPA_TEXTO_A_NUMERO = {
    "HOSPITALIZADO": 1, "HOSPITAL": 1, "NO SE ENCUENTRA EN EL DOMICILIO": 2,
    "NO SE EN CUENTRA EN EL DOMICILIO": 2, "NO ESTA": 2, "CITA MEDICA": 3,
    "CITA": 3, "SIN SERVICIO": 4, "NO GENERO RESIDUOS": 5, "NO TENIA": 5,
    "NO TIENE": 5, "BOTO LOS RESIDUOS A LA BASURA": 6, "ARROJO LOS RESIDUOS A LA BASURA": 6,
    "FALLECIDO": 7, "FALLECIO": 7, "DAÑOS DEL VEHICULO": 8, "FACTORES AMBIENTALES": 9,
    "LABORES ADMINISTRATIVAS": 10, "ALTERACION ORDEN PUBLICO": 11,
    "NO REQUIERE RECOLECCION": 12, "SERVICIOS SUSPENDIDOS TEMPORALMENTE": 13,
    "RECOLECCION CADA 15 DIAS": 14, "SE CAMBIO DOMICILIO": 15, "SE CAMBIO DE DOMICILIO": 15
}

def convertir_novedad(texto_nov):
    if pd.isna(texto_nov): return ""
    t = str(texto_nov).strip().upper()
    if "SIN NOVEDAD" in t or t == "": return ""
    for clave, num in MAPA_TEXTO_A_NUMERO.items():
        if clave in t: return num
    return ""

def extraer_catalogo_historico(archivo_bytes):
    catalogo = {}
    wb = openpyxl.load_workbook(archivo_bytes, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        col_ced, col_tipo, f_inicio = None, None, 1
        for r in range(1, min(25, ws.max_row + 1)):
            vals = [str(ws.cell(r, c).value or '').upper() for c in range(1, min(20, ws.max_column + 1))]
            for idx, val in enumerate(vals, start=1):
                if "CEDULA" in val or "DOCUMENTO" in val: col_ced = idx
                if "TIPO" in val and "PACIENTE" in val: col_tipo = idx
            if col_ced and col_tipo:
                f_inicio = r + 1
                break
        if col_ced and col_tipo:
            for r in range(f_inicio, ws.max_row + 1):
                c_val = ws.cell(r, col_ced).value
                t_val = ws.cell(r, col_tipo).value
                if c_val is not None and t_val is not None:
                    ced = str(c_val).split('.')[0].strip()
                    tipo = str(t_val).strip().upper()
                    if "CRONIC" in tipo: catalogo[ced] = "CRONICO"
                    elif "AGUD" in tipo: catalogo[ced] = "AGUDO"
    return catalogo

def cargar_excel_verificado(archivo_bytes):
    wb = openpyxl.load_workbook(archivo_bytes, data_only=True)
    sheet_name = 'CONSOLIDADO' if 'CONSOLIDADO' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    header_row = None
    headers = []
    for r in range(1, min(30, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or '').strip().upper() for c in range(1, min(20, ws.max_column + 1))]
        if any("CEDULA" in v or "DOCUMENTO" in v for v in vals) and any("EFECTIVA" in v or "RECOLECCION" in v for v in vals):
            header_row = r
            headers = [str(ws.cell(r, c).value or '').strip() for c in range(1, min(20, ws.max_column + 1))]
            break

    if not header_row:
        raise ValueError("No se encontró la tabla operativa en la hoja 'CONSOLIDADO'. Verifica las columnas CÉDULA y RECOLECCIÓN EFECTIVA.")

    data = []
    for r in range(header_row + 1, ws.max_row + 1):
        c_ced = ws.cell(r, 3).value
        c_nom = str(ws.cell(r, 4).value or '').upper()
        if "TOTAL" in c_nom: continue
        if c_ced is not None and str(c_ced).strip() != '':
            data.append([ws.cell(r, c).value for c in range(1, len(headers) + 1)])

    return pd.DataFrame(data, columns=headers)

# ==========================================
# 2. CONSTRUCCIÓN DE LA HOJA INFORMACION (PLANTILLA IDÉNTICA)
# ==========================================
def construir_hoja_informacion(ws_i, df_operativa, codigos_rf):
    ws_i.views.sheetView[0].showGridLines = True
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_reg = Font(name="Calibri", size=11)
    
    col_efe = [c for c in df_operativa.columns if 'EFECTIVA' in c.upper()][0]
    col_tip = [c for c in df_operativa.columns if 'TIPO' in c.upper()][0]
    col_nc = [c for c in df_operativa.columns if 'CRONIC' in c.upper() and 'NOVEDAD' in c.upper()]
    col_na = [c for c in df_operativa.columns if 'AGUD' in c.upper() and 'NOVEDAD' in c.upper()]
    col_ent = [c for c in df_operativa.columns if 'ENTIDAD' in c.upper()][0]

    # Bloque 1: Recoleccion Efectiva (A2:B5)
    ws_i.cell(2, 1, "SI/NO").font = font_bold
    ws_i.cell(2, 2, "RECOLECCION EFECTIVA").font = font_bold
    
    si_count = int((df_operativa[col_efe].astype(str).str.strip().str.upper() == 'SI').sum())
    no_count = int((df_operativa[col_efe].astype(str).str.strip().str.upper() == 'NO').sum())
    tot_visitas = len(df_operativa)
    
    ws_i.cell(3, 1, "SI").font = font_reg
    ws_i.cell(3, 2, si_count).font = font_reg
    ws_i.cell(4, 1, "NO").font = font_reg
    ws_i.cell(4, 2, no_count).font = font_reg
    ws_i.cell(5, 1, "TOTAL").font = font_bold
    ws_i.cell(5, 2, tot_visitas).font = font_bold
    
    # Bloque 2: Tipo de Paciente (D2:E5)
    ws_i.cell(2, 4, "TIPO PACIENTE").font = font_bold
    ws_i.cell(2, 5, "CANTIDAD DE RECOLECCION PROGRAMADA").font = font_bold
    
    cron_count = int((df_operativa[col_tip].astype(str).str.upper().str.contains('CRONIC')).sum())
    agud_count = int((df_operativa[col_tip].astype(str).str.upper().str.contains('AGUD')).sum())
    
    ws_i.cell(3, 4, "CRONICO").font = font_reg
    ws_i.cell(3, 5, cron_count).font = font_reg
    ws_i.cell(4, 4, "AGUDO").font = font_reg
    ws_i.cell(4, 5, agud_count).font = font_reg
    ws_i.cell(5, 4, "TOTAL").font = font_bold
    ws_i.cell(5, 5, cron_count + agud_count).font = font_bold
    
    # Bloque 3: Novedades Cronicos (A8:B15)
    ws_i.cell(8, 1, "TIPO DE NOVEDAD").font = font_bold
    ws_i.cell(8, 2, "CANTIDAD NOVEDAD CRONICO").font = font_bold
    
    r_c = 9
    tot_c = 0
    if col_nc:
        s_cron = df_operativa[df_operativa[col_nc[0]].notna() & (df_operativa[col_nc[0]].astype(str).str.strip() != '')][col_nc[0]].value_counts()
        for cod_num, cnt in s_cron.items():
            nom_nov = codigos_rf.get(int(float(cod_num)) if str(cod_num).replace('.','',1).isdigit() else 0, str(cod_num))
            ws_i.cell(r_c, 1, nom_nov).font = font_reg
            ws_i.cell(r_c, 2, int(cnt)).font = font_reg
            tot_c += int(cnt)
            r_c += 1
    ws_i.cell(r_c, 1, "TOTAL").font = font_bold
    ws_i.cell(r_c, 2, tot_c).font = font_bold
    
    # Bloque 4: Novedades Agudos (D8:E11)
    ws_i.cell(8, 4, "TIPO DE NOVEDAD").font = font_bold
    ws_i.cell(8, 5, "CANTIDAD NOVEDAD AGUDOS").font = font_bold
    
    r_a = 9
    tot_a = 0
    if col_na:
        s_agud = df_operativa[df_operativa[col_na[0]].notna() & (df_operativa[col_na[0]].astype(str).str.strip() != '')][col_na[0]].value_counts()
        for cod_num, cnt in s_agud.items():
            nom_nov = codigos_rf.get(int(float(cod_num)) if str(cod_num).replace('.','',1).isdigit() else 0, str(cod_num))
            ws_i.cell(r_a, 4, nom_nov).font = font_reg
            ws_i.cell(r_a, 5, int(cnt)).font = font_reg
            tot_a += int(cnt)
            r_a += 1
    ws_i.cell(r_a, 4, "TOTAL").font = font_bold
    ws_i.cell(r_a, 5, tot_a).font = font_bold
    
    # Bloque 5: Entidad (A17:B22)
    r_ent_start = max(max(r_c, r_a) + 2, 17)
    ws_i.cell(r_ent_start, 1, "ENTIDAD").font = font_bold
    ws_i.cell(r_ent_start, 2, "RECOLECCION POR ENTIDAD").font = font_bold
    
    s_ent = df_operativa[col_ent].value_counts()
    r_e = r_ent_start + 1
    for ent_nom, cnt in s_ent.items():
        ws_i.cell(r_e, 1, ent_nom).font = font_reg
        ws_i.cell(r_e, 2, int(cnt)).font = font_reg
        r_e += 1
    ws_i.cell(r_e, 1, "TOTAL").font = font_bold
    ws_i.cell(r_e, 2, int(s_ent.sum())).font = font_bold
    
    ws_i.column_dimensions['A'].width = 38
    ws_i.column_dimensions['B'].width = 28
    ws_i.column_dimensions['C'].width = 6
    ws_i.column_dimensions['D'].width = 38
    ws_i.column_dimensions['E'].width = 38

# ==========================================
# 3. GENERADORES DE ARCHIVOS EXCEL
# ==========================================
def generar_excel_fase1(df_operativa, codigos_rf):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        df_operativa.to_excel(writer, sheet_name='CONSOLIDADO', startrow=18, index=False)
        ws_c = writer.sheets['CONSOLIDADO']
        ws_c.views.sheetView[0].showGridLines = True
        
        ws_c.cell(1, 9, "CUADRO DE CLASIFICACION DOCUMENTAL").font = Font(bold=True)
        ws_c.cell(2, 9, "SIGLA DEPENDENCIA").font = Font(bold=True)
        ws_c.cell(2, 10, "CODIGO SERIE").font = Font(bold=True)
        ws_c.cell(2, 12, "ASUNTO").font = Font(bold=True)
        for cod, desc in codigos_rf.items():
            ws_c.cell(2 + cod, 9, "RF")
            ws_c.cell(2 + cod, 10, cod)
            ws_c.cell(2 + cod, 12, desc)
            
        fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        border_light = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                              top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

        for col in range(1, 16):
            c = ws_c.cell(row=19, column=col)
            c.fill = fill_header; c.font = font_header; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r in range(20, ws_c.max_row + 1):
            for col in range(1, 16):
                cell = ws_c.cell(row=r, column=col)
                cell.font = Font(name="Calibri", size=9)
                cell.border = border_light
                if col in [1, 2, 3, 10, 11, 12, 13, 14]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col in [6, 7, 8, 9]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if col in [7, 9]: cell.number_format = '#,##0.00'

        for col in ws_c.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            ws_c.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 11)

        ws_i = writer.book.create_sheet('INFORMACION')
        construir_hoja_informacion(ws_i, df_operativa, codigos_rf)

    out.seek(0)
    return out

def generar_excel_fase2_completo(df_ver, codigos_rf):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    azul_oscuro = "1F4E79"
    azul_claro = "D9E1F2"
    gris_fondo = "F2F2F2"
    verde_suave = "E2EFDA"
    rojo_suave = "FCE4D6"
    amarillo_suave = "FFF2CC"

    fill_hdr = PatternFill(start_color=azul_oscuro, end_color=azul_oscuro, fill_type="solid")
    fill_sub = PatternFill(start_color=azul_claro, end_color=azul_claro, fill_type="solid")
    fill_tot = PatternFill(start_color=gris_fondo, end_color=gris_fondo, fill_type="solid")

    font_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=9, bold=True, color="000000")
    font_bold = Font(name="Calibri", size=9, bold=True)
    font_regular = Font(name="Calibri", size=9)
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    # Hoja 1: CONSOLIDADO
    ws_c = wb.create_sheet('CONSOLIDADO')
    ws_c.views.sheetView[0].showGridLines = True

    ws_c.cell(1, 9, "CUADRO DE CLASIFICACION DOCUMENTAL").font = Font(bold=True)
    ws_c.cell(2, 9, "SIGLA DEPENDENCIA").font = Font(bold=True)
    ws_c.cell(2, 10, "CODIGO SERIE").font = Font(bold=True)
    ws_c.cell(2, 12, "ASUNTO").font = Font(bold=True)
    for cod, desc in codigos_rf.items():
        ws_c.cell(2 + cod, 9, "RF")
        ws_c.cell(2 + cod, 10, cod)
        ws_c.cell(2 + cod, 12, desc)

    cols_data = list(df_ver.columns)
    col_map = {}
    for c_i, col_name in enumerate(cols_data, start=1):
        col_letter = get_column_letter(c_i)
        c_up = col_name.upper()
        if 'CEDULA' in c_up or 'DOCUMENTO' in c_up: col_map['CEDULA'] = col_letter
        if 'NOMBRE' in c_up or 'USUARIO' in c_up: col_map['PACIENTE'] = col_letter
        if 'ENTIDAD' in c_up: col_map['ENTIDAD'] = col_letter
        if 'EFECTIVA' in c_up: col_map['EFECTIVA'] = col_letter
        if 'TIPO' in c_up and 'PACIENTE' in c_up: col_map['TIPO'] = col_letter
        if 'RESIDUOS GR' in c_up: col_map['RESIDUOS_GR'] = col_letter
        if 'PESO GR' in c_up: col_map['GUARDIAN_GR'] = col_letter
        if 'CRONIC' in c_up and 'NOVEDAD' in c_up: col_map['NOV_C'] = col_letter
        if 'AGUD' in c_up and 'NOVEDAD' in c_up: col_map['NOV_A'] = col_letter
        if 'FECHA' in c_up: col_map['FECHA'] = col_letter
        if 'OBSERVACION' in c_up: col_map['OBS'] = col_letter

        c = ws_c.cell(19, c_i, col_name)
        c.fill = fill_hdr
        c.font = font_hdr
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in df_ver.reset_index(drop=True).iterrows():
        row_num = r_idx + 20
        for c_i, col_name in enumerate(cols_data, start=1):
            val = row[col_name]
            cell = ws_c.cell(row_num, c_i, val)
            cell.font = font_regular
            cell.border = border_thin
            if str(val).replace('.', '', 1).isdigit():
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws_c.columns:
        max_l = max(len(str(c.value or '')) for c in col)
        ws_c.column_dimensions[get_column_letter(col[0].column)].width = max(max_l + 3, 11)

    max_data_r = len(df_ver) + 19
    l_efe = col_map.get('EFECTIVA', 'J')
    l_ced = col_map.get('CEDULA', 'C')
    l_res = col_map.get('RESIDUOS_GR', 'G')
    l_gua = col_map.get('GUARDIAN_GR', 'I')
    l_nc = col_map.get('NOV_C', 'L')
    l_na = col_map.get('NOV_A', 'M')

    # Hoja 2: INFORMACION (Plantilla idéntica a tu formato oficial)
    ws_i = wb.create_sheet('INFORMACION')
    construir_hoja_informacion(ws_i, df_ver, codigos_rf)

    # Hoja 3: TABLERO_INDICADORES_3100
    ws_k = wb.create_sheet('TABLERO_INDICADORES_3100')
    ws_k.views.sheetView[0].showGridLines = True
    ws_k.merge_cells("A1:I1")
    c1 = ws_k.cell(1, 1, "SISTEMA OBLIGATORIO DE GARANTÍA DE CALIDAD - RESOLUCIÓN 3100 DE 2019")
    c1.fill = fill_hdr; c1.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF"); c1.alignment = Alignment(horizontal="center", vertical="center")

    ws_k.merge_cells("A2:I2")
    c2 = ws_k.cell(2, 1, "TABLERO MAESTRO DE INDICADORES DE CALIDAD, BIORRIESGO Y TRAZABILIDAD EXTRAMURAL")
    c2.fill = fill_sub; c2.font = Font(name="Calibri", size=10, bold=True, color="1F4E79"); c2.alignment = Alignment(horizontal="center", vertical="center")

    hdrs_kpi = ["ITEM", "ESTÁNDAR RES. 3100", "NOMBRE DEL INDICADOR", "NUMERADOR (FÓRMULA)", "DENOMINADOR (FÓRMULA)", "FÓRMULA EXCEL", "META", "RESULTADO", "ESTADO"]
    ws_k.row_dimensions[4].row_height = 28
    for c_i, h in enumerate(hdrs_kpi, start=1):
        cell = ws_k.cell(4, c_i, h)
        cell.fill = fill_hdr; cell.font = font_hdr; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border_thin

    indicadores = [
        (1, "Procesos Prioritarios", "Adherencia a Ruta y Oportunidad Extramural",
         f'=COUNTIF(CONSOLIDADO!${l_efe}$20:${l_efe}${max_data_r}, "SI")',
         f'=COUNTA(CONSOLIDADO!${l_ced}$20:${l_ced}${max_data_r})',
         '=D5/E5', ">= 90%", "0.0%", '=IF(H5>=0.9, "CUMPLE", "ALERTA")'),
        (2, "Seguridad del Paciente", "Tasa de Eventos Críticos: Descarte en Basura Común (RF 6)",
         f'=COUNTIF(CONSOLIDADO!${l_nc}$20:${l_nc}${max_data_r}, 6)+COUNTIF(CONSOLIDADO!${l_na}$20:${l_na}${max_data_r}, 6)',
         f'=COUNTA(CONSOLIDADO!${l_ced}$20:${l_ced}${max_data_r})',
         '=D6/E6', "0.0%", "0.0%", '=IF(H6=0, "CUMPLE", "CRÍTICO")'),
        (3, "Infraestructura / PGIRASA", "Generación Neta Mensual de Residuos Biosanitarios (RH1)",
         f'=SUM(CONSOLIDADO!${l_res}$20:${l_res}${max_data_r})',
         "1000 (Gramos por Kg)",
         '=D7/1000', "Conciliado 100%", '#,##0.0 "Kg"', '"CONFORME"'),
        (4, "Infraestructura / PGIRASA", "Generación Neta Mensual de Cortopunzantes (Guardianes)",
         f'=SUM(CONSOLIDADO!${l_gua}$20:${l_gua}${max_data_r})',
         "1000 (Gramos por Kg)",
         '=D8/1000', "Conciliado 100%", '#,##0.0 "Kg"', '"CONFORME"'),
        (5, "Gestión Ambiental", "Intensidad de Generación Promedio por Atención Efectiva",
         f'=SUM(CONSOLIDADO!${l_res}$20:${l_res}${max_data_r})',
         f'=COUNTIF(CONSOLIDADO!${l_efe}$20:${l_efe}${max_data_r}, "SI")',
         '=D9/E9', "Estándar Técnico", '#,##0.0 "g/visita"', '"CONTROLADO"'),
        (6, "Modalidad Domiciliaria", "Cumplimiento Frecuencia Semanal Crónicos (>= 4 Visitas/Mes)",
         '=COUNTIF(TRAZABILIDAD_CRONICOS!E15:E1000, ">=4")',
         '=COUNTA(TRAZABILIDAD_CRONICOS!A15:A1000)',
         '=D10/E10', ">= 85%", "0.0%", '=IF(H10>=0.85, "CUMPLE", "ALERTA")')
    ]

    for idx, (item, est, nom, num_f, den_f, res_f, meta, fmt, est_f) in enumerate(indicadores, start=5):
        ws_k.row_dimensions[idx].row_height = 24
        ws_k.cell(idx, 1, item).alignment = Alignment(horizontal="center", vertical="center")
        ws_k.cell(idx, 2, est).alignment = Alignment(horizontal="left", vertical="center")
        ws_k.cell(idx, 3, nom).alignment = Alignment(horizontal="left", vertical="center")
        
        c_num = ws_k.cell(idx, 4, num_f); c_num.alignment = Alignment(horizontal="right", vertical="center")
        if num_f.startswith('='): c_num.number_format = '#,##0'
        c_den = ws_k.cell(idx, 5, den_f); c_den.alignment = Alignment(horizontal="right", vertical="center")
        if str(den_f).startswith('='): c_den.number_format = '#,##0'

        c_res = ws_k.cell(idx, 6, res_f); c_res.alignment = Alignment(horizontal="right", vertical="center"); c_res.number_format = fmt; c_res.font = font_bold
        c_meta = ws_k.cell(idx, 7, meta); c_meta.alignment = Alignment(horizontal="center", vertical="center"); c_meta.font = font_bold
        c_res2 = ws_k.cell(idx, 8, f'=F{idx}'); c_res2.alignment = Alignment(horizontal="right", vertical="center"); c_res2.number_format = fmt; c_res2.font = font_bold
        c_est = ws_k.cell(idx, 9, est_f); c_est.alignment = Alignment(horizontal="center", vertical="center"); c_est.font = font_bold

        for c_i in range(1, 10):
            ws_k.cell(idx, c_i).border = border_thin
            if c_i not in [1, 6, 7, 8, 9]: ws_k.cell(idx, c_i).font = font_regular

    ws_k.column_dimensions['A'].width = 6
    ws_k.column_dimensions['B'].width = 24
    ws_k.column_dimensions['C'].width = 38
    ws_k.column_dimensions['D'].width = 28
    ws_k.column_dimensions['E'].width = 28
    ws_k.column_dimensions['F'].width = 16
    ws_k.column_dimensions['G'].width = 18
    ws_k.column_dimensions['H'].width = 16
    ws_k.column_dimensions['I'].width = 14

    # Hoja 4: TRAZABILIDAD_CRONICOS
    ws_cr = wb.create_sheet('TRAZABILIDAD_CRONICOS')
    ws_cr.views.sheetView[0].showGridLines = True
    ws_cr.merge_cells("A1:G1")
    c1 = ws_cr.cell(1, 1, "TRAZABILIDAD Y AUDITORÍA DE FRECUENCIA EN PACIENTES CRÓNICOS (1 VISITA / SEMANA)")
    c1.fill = fill_hdr; c1.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF"); c1.alignment = Alignment(horizontal="center", vertical="center")

    ws_cr.merge_cells("A3:E3")
    c3 = ws_cr.cell(3, 1, "DISTRIBUCIÓN GENERAL DE CUMPLIMIENTO MENSUAL (META ≥ 4 VISITAS)")
    c3.fill = fill_sub; c3.font = Font(name="Calibri", size=9, bold=True, color="1F4E79"); c3.alignment = Alignment(horizontal="center", vertical="center")

    headers_res = ["FRECUENCIA MENSUAL", "PACIENTES AUDITADOS", "% DEL TOTAL", "ESTADO DE CUMPLIMIENTO", "INTERPRETACIÓN CLÍNICA"]
    for c_i, h in enumerate(headers_res, start=1):
        cell = ws_cr.cell(4, c_i, h)
        cell.fill = fill_hdr; cell.font = font_hdr; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = border_thin

    col_ced_n = [c for c in df_ver.columns if 'CEDULA' in c.upper()][0]
    col_nom_n = [c for c in df_ver.columns if 'NOMBRE' in c.upper() or 'USUARIO' in c.upper()][0]
    col_ent_n = [c for c in df_ver.columns if 'ENTIDAD' in c.upper()][0]
    col_tip_n = [c for c in df_ver.columns if 'TIPO' in c.upper()][0]

    df_cron = df_ver[df_ver[col_tip_n].astype(str).str.upper().str.contains('CRONIC')].copy()
    pacientes_unicos = df_cron[[col_ced_n, col_nom_n, col_ent_n]].drop_duplicates().reset_index(drop=True)
    tot_pacs = len(pacientes_unicos)
    f_ini_det = 15
    f_fin_det = f_ini_det + tot_pacs - 1

    niveles = [
        (">= 5 Recolecciones", f'=COUNTIF(E{f_ini_det}:E{f_fin_det}, ">=5")', "Óptimo", "Cobertura completa en mes de 5 semanas", verde_suave),
        ("4 Recolecciones", f'=COUNTIF(E{f_ini_det}:E{f_fin_det}, 4)', "Meta Cumplida", "1 recolección semanal garantizada (Meta Res. 3100)", verde_suave),
        ("3 Recolecciones", f'=COUNTIF(E{f_ini_det}:E{f_fin_det}, 3)', "Alerta Moderada", "Falla en 1 semana (acumulación de biorriesgo)", amarillo_suave),
        ("<= 2 Recolecciones", f'=COUNTIF(E{f_ini_det}:E{f_fin_det}, "<=2")', "Riesgo Crítico", "Falla reiterada (residuos acumulados >= 14 días)", rojo_suave)
    ]

    for idx_n, (label, form_p, est, interp, bg_col) in enumerate(niveles, start=5):
        ws_cr.row_dimensions[idx_n].row_height = 20
        ws_cr.cell(idx_n, 1, label).alignment = Alignment(horizontal="left", vertical="center")
        c_p = ws_cr.cell(idx_n, 2, form_p); c_p.alignment = Alignment(horizontal="right", vertical="center"); c_p.font = font_bold
        c_pct = ws_cr.cell(idx_n, 3, f'=B{idx_n}/B9'); c_pct.alignment = Alignment(horizontal="right", vertical="center"); c_pct.number_format = "0.0%"; c_pct.font = font_bold
        c_est = ws_cr.cell(idx_n, 4, est); c_est.alignment = Alignment(horizontal="center", vertical="center"); c_est.font = font_bold
        c_est.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
        ws_cr.cell(idx_n, 5, interp).alignment = Alignment(horizontal="left", vertical="center")
        for c_i in range(1, 6):
            ws_cr.cell(idx_n, c_i).border = border_thin
            if c_i != 4: ws_cr.cell(idx_n, c_i).font = font_regular if c_i not in [2, 3] else font_bold

    ws_cr.row_dimensions[9].row_height = 20
    ws_cr.cell(9, 1, "TOTAL PACIENTES CRÓNICOS").alignment = Alignment(horizontal="left", vertical="center")
    ws_cr.cell(9, 2, '=SUM(B5:B8)').alignment = Alignment(horizontal="right", vertical="center")
    ws_cr.cell(9, 3, '=SUM(C5:C8)').alignment = Alignment(horizontal="right", vertical="center")
    ws_cr.cell(9, 3).number_format = "0.0%"
    ws_cr.cell(9, 4, "100% AUDITADO").alignment = Alignment(horizontal="center", vertical="center")
    ws_cr.cell(9, 5, "Población total con indicación de recolección semanal").alignment = Alignment(horizontal="left", vertical="center")
    for c_i in range(1, 6):
        c = ws_cr.cell(9, c_i); c.border = border_thin; c.fill = fill_tot; c.font = font_bold

    # Detalle paciente por paciente
    ws_cr.merge_cells("A12:G12")
    c12 = ws_cr.cell(12, 1, "AUDITORÍA NOMINAL: DETALLE DE RECOLECCIONES EFECTIVAS POR USUARIO")
    c12.fill = fill_hdr; c12.font = font_hdr; c12.alignment = Alignment(horizontal="center", vertical="center")

    headers_det = ["CÉDULA", "NOMBRE DEL USUARIO", "ENTIDAD", "VISITAS PROGRAMADAS", "RECOLECCIONES EFECTIVAS (SI)", "NOVEDADES (NO)", "ESTADO DE COBERTURA"]
    ws_cr.row_dimensions[14].row_height = 24
    for c_i, h in enumerate(headers_det, start=1):
        cell = ws_cr.cell(14, c_i, h); cell.fill = fill_sub; cell.font = font_sub; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = border_thin

    for idx_p, r_pac in pacientes_unicos.iterrows():
        fila_r = f_ini_det + idx_p
        ced_val = str(r_pac[col_ced_n]).split('.')[0].strip()
        nom_val = str(r_pac[col_nom_n]).strip().upper()
        ent_val = str(r_pac[col_ent_n]).strip()

        ws_cr.row_dimensions[fila_r].row_height = 18
        ws_cr.cell(fila_r, 1, ced_val).alignment = Alignment(horizontal="center", vertical="center")
        ws_cr.cell(fila_r, 2, nom_val).alignment = Alignment(horizontal="left", vertical="center")
        ws_cr.cell(fila_r, 3, ent_val).alignment = Alignment(horizontal="left", vertical="center")

        c_prog = ws_cr.cell(fila_r, 4, f'=COUNTIF(CONSOLIDADO!${l_ced}$20:${l_ced}${max_data_r}, A{fila_r})'); c_prog.alignment = Alignment(horizontal="right", vertical="center")
        c_eff = ws_cr.cell(fila_r, 5, f'=COUNTIFS(CONSOLIDADO!${l_ced}$20:${l_ced}${max_data_r}, A{fila_r}, CONSOLIDADO!${l_efe}$20:${l_efe}${max_data_r}, "SI")'); c_eff.alignment = Alignment(horizontal="right", vertical="center"); c_eff.font = font_bold
        c_nov = ws_cr.cell(fila_r, 6, f'=COUNTIFS(CONSOLIDADO!${l_ced}$20:${l_ced}${max_data_r}, A{fila_r}, CONSOLIDADO!${l_efe}$20:${l_efe}${max_data_r}, "NO")'); c_nov.alignment = Alignment(horizontal="right", vertical="center")
        c_est = ws_cr.cell(fila_r, 7, f'=IF(E{fila_r}>=4, "CUMPLE SEMANAL", "ALERTA RIESGO")'); c_est.alignment = Alignment(horizontal="center", vertical="center"); c_est.font = font_bold

        for c_i in range(1, 8):
            ws_cr.cell(fila_r, c_i).border = border_thin
            if c_i not in [5, 7]: ws_cr.cell(fila_r, c_i).font = font_regular

    ws_cr.column_dimensions['A'].width = 16
    ws_cr.column_dimensions['B'].width = 36
    ws_cr.column_dimensions['C'].width = 18
    ws_cr.column_dimensions['D'].width = 22
    ws_cr.column_dimensions['E'].width = 26
    ws_cr.column_dimensions['F'].width = 18
    ws_cr.column_dimensions['G'].width = 22

    # Hoja 5: DETALLE_NOVEDADES
    ws_n = wb.create_sheet('DETALLE_NOVEDADES')
    ws_n.views.sheetView[0].showGridLines = True
    ws_n.merge_cells("A1:I1")
    c1 = ws_n.cell(1, 1, "REGISTRO NOMINAL DE NOVEDADES OPERATIVAS (VISITAS NO EFECTIVAS - RF 1 A 15)")
    c1.fill = fill_hdr; c1.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF"); c1.alignment = Alignment(horizontal="center", vertical="center")

    headers_nov = ["ITEM", "FECHA", "CÉDULA", "NOMBRE DEL USUARIO", "ENTIDAD", "TIPO PACIENTE", "CÓDIGO RF", "DESCRIPCIÓN OFICIAL NOVEDAD", "OBSERVACIÓN EN CAMPO"]
    ws_n.row_dimensions[3].row_height = 24
    for c_i, h in enumerate(headers_nov, start=1):
        cell = ws_n.cell(3, c_i, h); cell.fill = fill_hdr; cell.font = font_hdr; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = border_thin

    col_efe_n = [c for c in df_ver.columns if 'EFECTIVA' in c.upper()][0]
    col_fec_n = [c for c in df_ver.columns if 'FECHA' in c.upper()][0]
    col_obs_n = [c for c in df_ver.columns if 'OBSERVACION' in c.upper()][0]
    col_nc = [c for c in df_ver.columns if 'CRONIC' in c.upper() and 'NOVEDAD' in c.upper()]
    col_na = [c for c in df_ver.columns if 'AGUD' in c.upper() and 'NOVEDAD' in c.upper()]

    df_nov = df_ver[df_ver[col_efe_n].astype(str).str.strip().str.upper() == 'NO'].copy()

    for idx_n, (_, row) in enumerate(df_nov.iterrows(), start=1):
        fila_r = 3 + idx_n
        ws_n.row_dimensions[fila_r].row_height = 18
        cod_val = ""
        if col_nc and pd.notna(row[col_nc[0]]) and str(row[col_nc[0]]).strip() != '':
            cod_val = str(row[col_nc[0]]).split('.')[0].strip()
        elif col_na and pd.notna(row[col_na[0]]) and str(row[col_na[0]]).strip() != '':
            cod_val = str(row[col_na[0]]).split('.')[0].strip()

        desc_nov = codigos_rf.get(int(float(cod_val)) if str(cod_val).replace('.','',1).isdigit() else 0, "NOVEDAD REGISTRADA")
        fec_val = str(row[col_fec_n]).split(' ')[0]
        ced_val = str(row[col_ced_n]).split('.')[0].strip()
        nom_val = str(row[col_nom_n]).strip().upper()
        ent_val = str(row[col_ent_n]).strip()
        tip_val = str(row[col_tip_n]).strip().upper()
        obs_val = str(row[col_obs_n] or '').strip()
        if obs_val in ['nan', 'None']: obs_val = ''

        ws_n.cell(fila_r, 1, idx_n).alignment = Alignment(horizontal="center", vertical="center")
        ws_n.cell(fila_r, 2, fec_val).alignment = Alignment(horizontal="center", vertical="center")
        ws_n.cell(fila_r, 3, ced_val).alignment = Alignment(horizontal="center", vertical="center")
        ws_n.cell(fila_r, 4, nom_val).alignment = Alignment(horizontal="left", vertical="center")
        ws_n.cell(fila_r, 5, ent_val).alignment = Alignment(horizontal="left", vertical="center")
        ws_n.cell(fila_r, 6, tip_val).alignment = Alignment(horizontal="center", vertical="center")
        ws_n.cell(fila_r, 7, f"RF-{cod_val}").alignment = Alignment(horizontal="center", vertical="center")
        ws_n.cell(fila_r, 8, desc_nov).alignment = Alignment(horizontal="left", vertical="center")
        ws_n.cell(fila_r, 9, obs_val).alignment = Alignment(horizontal="left", vertical="center")

        for c_i in range(1, 10):
            ws_n.cell(fila_r, c_i).border = border_thin
            if c_i not in [1, 7]: ws_n.cell(fila_r, c_i).font = font_regular

    ws_n.column_dimensions['A'].width = 6
    ws_n.column_dimensions['B'].width = 14
    ws_n.column_dimensions['C'].width = 16
    ws_n.column_dimensions['D'].width = 36
    ws_n.column_dimensions['E'].width = 18
    ws_n.column_dimensions['F'].width = 16
    ws_n.column_dimensions['G'].width = 14
    ws_n.column_dimensions['H'].width = 34
    ws_n.column_dimensions['I'].width = 36

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

# ==========================================
# 4. INTERFAZ Y FLUJO DE TRABAJO
# ==========================================
fase = st.sidebar.radio(
    "Selecciona la acción:",
    ["1️⃣ Cruce Inicial (Generar Preliminar)", "2️⃣ Cargar Archivo Verificado (Indicadores Finales)"]
)

if fase == "1️⃣ Cruce Inicial (Generar Preliminar)":
    st.subheader("Fase 1: Cruce de Archivo Web vs. Histórico")
    st.info("Sube las visitas del sistema web y tu histórico para generar el libro Excel oficial (`CONSOLIDADO` e `INFORMACION`) para revisión.")

    c1, c2 = st.columns(2)
    with c1:
        f_web = st.file_uploader("📥 Visitas del sistema web (.xls / .xlsx)", type=["xls", "xlsx"], key="web_raw")
    with c2:
        f_hist = st.file_uploader("📚 Archivo histórico / mes anterior (.xlsx)", type=["xlsx"], key="hist_raw")

    if f_web and f_hist:
        if st.button("⚙️ Realizar Cruce y Generar Archivo Oficial", type="primary"):
            with st.spinner("Procesando datos, cruzando pacientes y estructurando hojas..."):
                try:
                    try: df_v = pd.read_html(f_web, encoding='utf-8')[0]
                    except Exception: df_v = pd.read_excel(f_web)

                    cat = extraer_catalogo_historico(f_hist)
                    df_v['Cedula_Str'] = df_v['Cédula'].astype(str).str.split('.').str[0].str.strip()
                    df_v['Fecha_DT'] = pd.to_datetime(df_v['Fecha'], format='%d/%m/%Y', errors='coerce')
                    df_v = df_v.sort_values(by=['Fecha_DT', 'Hora'], ascending=[True, True]).reset_index(drop=True)

                    df_v['Tipo_Final'] = df_v['Cedula_Str'].map(cat).fillna('')
                    df_v['Num_Nov'] = df_v['Novedad'].apply(convertir_novedad)
                    df_v['Efectiva'] = np.where(df_v['Num_Nov'] != '', 'NO', 'SI')

                    df_v['Nov_C'] = np.where((df_v['Tipo_Final'] == 'CRONICO') & (df_v['Num_Nov'] != ''), df_v['Num_Nov'], '')
                    df_v['Nov_A'] = np.where((df_v['Tipo_Final'] == 'AGUDO') & (df_v['Num_Nov'] != ''), df_v['Num_Nov'], '')
                    mascara_sin = (df_v['Tipo_Final'] == '') & (df_v['Num_Nov'] != '')
                    df_v.loc[mascara_sin, 'Nov_C'] = df_v.loc[mascara_sin, 'Num_Nov']

                    iso_c = df_v['Fecha_DT'].dt.isocalendar()
                    df_v['Sem_Ano'] = iso_c.year.astype(str) + "-W" + iso_c.week.astype(str)
                    df_v['Vis_Sem'] = df_v.groupby(['Cedula_Str', 'Sem_Ano']).cumcount() + 1
                    df_v['Repeticion'] = df_v['Vis_Sem'].apply(lambda x: 'X' if x > 1 else '')

                    df_op = pd.DataFrame({
                        'ITEM': range(1, len(df_v) + 1),
                        'FECHA': df_v['Fecha_DT'].dt.strftime('%d/%m/%Y'),
                        'CEDULA': df_v['Cedula_Str'],
                        'NOMBRE DEL USUARIO': df_v['Paciente'].str.strip().str.upper(),
                        'ENTIDAD': df_v['Entidad'].str.strip(),
                        'NUMERO DE BOLSAS': pd.to_numeric(df_v['Bolsas'], errors='coerce').fillna(0).astype(int),
                        'RESIDUOS GR RECOLECTADOS ': pd.to_numeric(df_v['Peso Residuos (g)'], errors='coerce').fillna(0.0),
                        'GUARDIAN': pd.to_numeric(df_v['Cant. Guardian'], errors='coerce').fillna(0).astype(int),
                        'PESO GR': pd.to_numeric(df_v['Peso Guardian (g)'], errors='coerce').fillna(0.0),
                        'RECOLECCION EFECTIVA': df_v['Efectiva'],
                        'TIPO DE PACIENTE ': df_v['Tipo_Final'],
                        'CANTIDAD NOVEDAD  CRONICOS': df_v['Nov_C'],
                        'CANTIDAD NOVEDAD AGUDOS': df_v['Nov_A'],
                        'REPETICION POR NOVEDAD': df_v['Repeticion'],
                        'OBSERVACION': df_v['Observaciones'].replace('Sin observaciones', '')
                    })

                    out_excel_f1 = generar_excel_fase1(df_op, CODIGOS_NOVEDADES_RF)
                    st.success("✅ Cruce preliminar realizado con éxito. Incluye las hojas `CONSOLIDADO` e `INFORMACION`:")
                    
                    st.download_button(
                        label="📥 Descargar Excel Oficial para Revisión (.xlsx)",
                        data=out_excel_f1,
                        file_name="CONSOLIDADO_RUTA_MENSUAL.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )
                    st.dataframe(df_op.head(15), use_container_width=True)

                except Exception as e:
                    st.error(f"Error procesando cruce preliminar: {e}")

else:
    st.subheader("Fase 2: Cargar Archivo Consolidado Verificado")
    st.markdown("Sube aquí el archivo Excel verificado para auditar los indicadores, analizar pacientes crónicos y descargar el libro completo.")

    f_verificado = st.file_uploader("📋 Sube el archivo Excel verificado (.xlsx)", type=["xlsx"], key="verificado_upload")

    if f_verificado:
        if st.button("📊 Generar Indicadores y Libro Completo de Auditoría", type="primary"):
            with st.spinner("Auditando datos verificados y calculando indicadores..."):
                try:
                    df_ver = cargar_excel_verificado(f_verificado)

                    col_ced = [c for c in df_ver.columns if 'CEDULA' in c.upper() or 'DOCUMENTO' in c.upper()][0]
                    col_pac = [c for c in df_ver.columns if 'NOMBRE' in c.upper() or 'USUARIO' in c.upper()][0]
                    col_ent = [c for c in df_ver.columns if 'ENTIDAD' in c.upper()][0]
                    col_efe = [c for c in df_ver.columns if 'EFECTIVA' in c.upper()][0]
                    col_tip = [c for c in df_ver.columns if 'TIPO' in c.upper()][0]
                    col_res = [c for c in df_ver.columns if 'RESIDUOS GR' in c.upper()][0]
                    col_gua = [c for c in df_ver.columns if 'PESO GR' in c.upper()][0]

                    col_nov_c_list = [c for c in df_ver.columns if 'CRONIC' in c.upper() and 'NOVEDAD' in c.upper()]
                    col_nov_a_list = [c for c in df_ver.columns if 'AGUD' in c.upper() and 'NOVEDAD' in c.upper()]
                    col_nc = col_nov_c_list[0] if col_nov_c_list else None
                    col_na = col_nov_a_list[0] if col_nov_a_list else None

                    df_ver['CED_STR'] = df_ver[col_ced].astype(str).str.split('.').str[0].str.strip()
                    df_ver['EFE_STR'] = df_ver[col_efe].astype(str).str.strip().str.upper()
                    df_ver['TIP_STR'] = df_ver[col_tip].astype(str).str.strip().str.upper()

                    def limpiar_num(val):
                        if pd.isna(val) or val is None or str(val).strip() == '': return 0.0
                        try: return float(str(val).replace(',', '').strip())
                        except: return 0.0

                    df_ver['PESO_RES'] = df_ver[col_res].apply(limpiar_num)
                    df_ver['PESO_GUA'] = df_ver[col_gua].apply(limpiar_num)

                    total_visitas = len(df_ver)
                    efectivas = (df_ver['EFE_STR'] == 'SI').sum()
                    novedades = (df_ver['EFE_STR'] == 'NO').sum()
                    adherencia = (efectivas / total_visitas * 100) if total_visitas > 0 else 0.0
                    tot_kg_bio = df_ver['PESO_RES'].sum() / 1000
                    tot_kg_gua = df_ver['PESO_GUA'].sum() / 1000
                    promedio_g = df_ver[df_ver['EFE_STR'] == 'SI']['PESO_RES'].mean()

                    def es_codigo_6(val):
                        if pd.isna(val): return False
                        s = str(val).split('.')[0].strip()
                        return s == '6' or 'BASURA' in str(val).upper()

                    disp_basura = 0
                    if col_nc: disp_basura += df_ver[col_nc].apply(es_codigo_6).sum()
                    if col_na: disp_basura += df_ver[col_na].apply(es_codigo_6).sum()

                    # Frecuencia en Crónicos
                    df_cron = df_ver[df_ver['TIP_STR'].str.contains('CRONIC')].copy()
                    pacientes_cron_unicos = df_cron[['CED_STR', col_pac, col_ent]].drop_duplicates().reset_index(drop=True)
                    tot_cron_unicos = len(pacientes_cron_unicos)

                    cron_efectivas = df_cron[df_cron['EFE_STR'] == 'SI'].groupby('CED_STR').size().reset_index(name='Recolecciones_Realizadas')
                    seguimiento_cron = pd.merge(pacientes_cron_unicos, cron_efectivas, on='CED_STR', how='left').fillna({'Recolecciones_Realizadas': 0})
                    seguimiento_cron['Recolecciones_Realizadas'] = seguimiento_cron['Recolecciones_Realizadas'].astype(int)

                    dist_cron = seguimiento_cron['Recolecciones_Realizadas'].value_counts().sort_index(ascending=False).reset_index()
                    dist_cron.columns = ['Recolecciones en el Mes', 'Cantidad de Pacientes']
                    dist_cron['Porcentaje del Total (%)'] = (dist_cron['Cantidad de Pacientes'] / tot_cron_unicos * 100).round(1)

                    # Generar el libro Excel completo de auditoría
                    excel_auditoria = generar_excel_fase2_completo(df_ver, CODIGOS_NOVEDADES_RF)

                    st.success(f"✅ Análisis completado: {total_visitas} visitas auditadas.")

                    # Botón destacado de descarga
                    st.download_button(
                        label="📥 Descargar Libro Completo con Fórmulas e Indicadores (.xlsx)",
                        data=excel_auditoria,
                        file_name="REPORTE_OFICIAL_INDICADORES_RES_3100.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )

                    tab1, tab2, tab3 = st.tabs([
                        "📊 Indicadores Calidad Res. 3100",
                        "👥 Trazabilidad Frecuencia Crónicos",
                        "📑 Vista Previa de Datos"
                    ])

                    with tab1:
                        st.subheader("Indicadores de Biorriesgo y RH1 (Resolución 3100)")
                        m1, m2, m3, m4 = st.columns(4)
                        m1.metric("Adherencia a Ruta (Meta ≥ 90%)", f"{adherencia:.1f}%")
                        m2.metric("Disposición en Basura (RF 6)", f"{disp_basura} casos", delta="Meta 0 casos", delta_color="inverse")
                        m3.metric("Biosanitarios (Kg para RH1)", f"{tot_kg_bio:,.2f} Kg")
                        m4.metric("Promedio g / Visita Efectiva", f"{promedio_g:.1f} g")

                        m5, m6, m7, m8 = st.columns(4)
                        m5.metric("Total Visitas Programadas", f"{total_visitas}")
                        m6.metric("Recolecciones Efectivas", f"{efectivas}")
                        m7.metric("Novedades Registradas", f"{novedades}")
                        m8.metric("Cortopunzantes (Guardianes)", f"{tot_kg_gua:,.2f} Kg")

                    with tab2:
                        st.subheader("Frecuencia de Recolección en Pacientes Crónicos (1 Visita / Semana)")
                        cumplimiento_ok = dist_cron[dist_cron['Recolecciones en el Mes'] >= 4]['Cantidad de Pacientes'].sum()
                        pct_ok = (cumplimiento_ok / tot_cron_unicos * 100) if tot_cron_unicos > 0 else 0

                        k1, k2, k3 = st.columns(3)
                        k1.metric("Pacientes Crónicos Únicos", f"{tot_cron_unicos}")
                        k2.metric("Cumplimiento Semanal (≥ 4 Visitas)", f"{cumplimiento_ok} pacientes", f"{pct_ok:.1f}% del total")
                        k3.metric("Pacientes en Riesgo (≤ 3 Visitas)", f"{tot_cron_unicos - cumplimiento_ok} pacientes", delta_color="inverse")

                        col_t, col_g = st.columns([1, 1])
                        with col_t:
                            st.dataframe(dist_cron, use_container_width=True, hide_index=True)
                        with col_g:
                            st.bar_chart(dist_cron.set_index('Recolecciones en el Mes')['Porcentaje del Total (%)'])

                    with tab3:
                        st.subheader("Datos Verificados de Entrada")
                        st.dataframe(df_ver.head(15), use_container_width=True)

                except Exception as e:
                    st.error(f"Error procesando el archivo verificado: {e}")